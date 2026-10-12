"""Excel / CSV helpers: read an uploaded sheet, write and stream files, convert cell values."""

import csv
import io
import json
import os
import tempfile
import zipfile
from collections.abc import Iterable, Iterator, Sequence
from datetime import date, datetime, time
from decimal import Decimal
from enum import Enum
from typing import Any, BinaryIO
from uuid import UUID

from fastapi.responses import FileResponse, Response, StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from starlette.background import BackgroundTask

from app.utils.exception_handler import CustomException, ExceptionType

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MEDIA_TYPE = "text/csv; charset=utf-8"
DATA_SHEET = "data"
# Excel hard limit is 1,048,576 rows per sheet (header included): continue on data_2, data_3...
SHEET_MAX_ROWS = 1_048_576
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="2F5597")
_REQUIRED_FILL = PatternFill("solid", fgColor="C00000")


def column_letter(index: int) -> str:
    """0-based column index -> Excel letter."""
    return get_column_letter(index + 1)


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def to_cell_value(value: Any) -> Any:
    """Python value -> value openpyxl can store. Lists/objects become JSON."""
    if value is None or isinstance(value, bool | int | float | Decimal | date | time):
        if isinstance(value, datetime) and value.tzinfo is not None:
            return value.isoformat()  # Excel has no timezone: keep it exact as text
        return value
    if isinstance(value, Enum):
        return to_cell_value(value.value)
    if isinstance(value, UUID):
        return str(value)
    if isinstance(value, list | tuple | set | dict):
        return json.dumps(value, ensure_ascii=False, default=str)
    return ILLEGAL_CHARACTERS_RE.sub("", str(value))


def _write_cell(sheet: Any, value: Any) -> WriteOnlyCell:
    cell = WriteOnlyCell(sheet, value=to_cell_value(value))
    if isinstance(cell.value, str) and cell.value.startswith(_FORMULA_PREFIXES):
        cell.data_type = "s"  # never let exported text run as a formula
    return cell


def _new_data_sheet(workbook: Workbook, headers: Sequence[str], number: int) -> Any:
    sheet = workbook.create_sheet(DATA_SHEET if number == 1 else f"{DATA_SHEET}_{number}")
    for index, header in enumerate(headers):
        sheet.column_dimensions[column_letter(index)].width = max(12, min(len(header) + 4, 40))
    header_cells = []
    for header in headers:
        cell = WriteOnlyCell(sheet, value=header)
        cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
        header_cells.append(cell)
    sheet.append(header_cells)
    return sheet


def write_data_workbook(target: Any, headers: Sequence[str], rows: Iterable[Sequence[Any]]) -> None:
    """Write-only workbook (rows go to temp files, memory stays flat) into a path or file."""
    workbook = Workbook(write_only=True)
    sheet_number, sheet_rows = 1, 1
    sheet = _new_data_sheet(workbook, headers, sheet_number)
    for row in rows:
        if sheet_rows >= SHEET_MAX_ROWS:
            sheet_number += 1
            sheet, sheet_rows = _new_data_sheet(workbook, headers, sheet_number), 1
        sheet.append([_write_cell(sheet, value) for value in row])
        sheet_rows += 1
    workbook.save(target)


def xlsx_file_response(
    headers: Sequence[str], rows: Iterable[Sequence[Any]], filename: str
) -> FileResponse:
    """Build the workbook in a temp file, stream it in chunks, delete it afterwards."""
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    try:
        write_data_workbook(path, headers, rows)
    except BaseException:
        os.unlink(path)
        raise
    return FileResponse(
        path,
        media_type=XLSX_MEDIA_TYPE,
        filename=filename,
        background=BackgroundTask(os.unlink, path),
    )


def _csv_value(value: Any) -> Any:
    value = to_cell_value(value)
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, date | time):
        return value.isoformat()
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return f"'{value}"  # OWASP CSV injection guard
    return value


def csv_streaming_response(
    headers: Sequence[str], rows: Iterable[Sequence[Any]], filename: str, chunk_rows: int = 500
) -> StreamingResponse:
    """True streaming: bytes leave as soon as rows are read. UTF-8 BOM so Excel shows accents."""

    def generate() -> Iterator[bytes]:
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        buffer.write("\ufeff")
        writer.writerow(headers)
        for count, row in enumerate(rows, start=1):
            writer.writerow([_csv_value(value) for value in row])
            if count % chunk_rows == 0:
                yield buffer.getvalue().encode()
                buffer.seek(0)
                buffer.truncate()
        yield buffer.getvalue().encode()

    return StreamingResponse(
        generate(),
        media_type=CSV_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def build_template_workbook(columns: Sequence[dict[str, Any]]) -> bytes:
    """Import template. columns: [{"name", "type", "required", "description"}]."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = DATA_SHEET
    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=column["name"])
        cell.font = _HEADER_FONT
        cell.fill = _REQUIRED_FILL if column["required"] else _HEADER_FILL
        note = f"{column['type']}{' - bắt buộc' if column['required'] else ''}"
        if column.get("description"):
            note += f"\n{column['description']}"
        cell.comment = Comment(note, "system")
        sheet.column_dimensions[get_column_letter(index)].width = max(14, len(column["name"]) + 4)
    sheet.freeze_panes = "A2"

    guide = workbook.create_sheet("guide")
    guide.append(["Cột", "Kiểu dữ liệu", "Bắt buộc", "Mô tả"])
    for cell in guide[1]:
        cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
    for column in columns:
        guide.append(
            [
                column["name"],
                column["type"],
                "x" if column["required"] else "",
                column.get("description") or "",
            ]
        )
    for letter, width in zip("ABCD", (24, 34, 10, 70), strict=True):
        guide.column_dimensions[letter].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def read_data_sheet(
    source: bytes | BinaryIO, max_rows: int | None
) -> tuple[list[Any], list[tuple[int, list[Any]]]]:
    """Return (header cells, [(excel row number, cells)]) of the "data" sheet (or the first one).
    Empty rows are skipped. source: bytes or a binary file (read from disk, not loaded in RAM).
    max_rows: None/0 = no limit."""
    if isinstance(source, bytes):
        source = io.BytesIO(source)
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as e:
        raise CustomException(ExceptionType.BAD_REQUEST, "Invalid Excel file (.xlsx)") from e
    try:
        sheet = (
            workbook[DATA_SHEET] if DATA_SHEET in workbook.sheetnames else workbook.worksheets[0]
        )
        rows = sheet.iter_rows(values_only=True)
        header = list(next(rows, None) or [])
        if not any(not is_blank(cell) for cell in header):
            raise CustomException(ExceptionType.BAD_REQUEST, "The first row must contain headers")
        data = []
        for row_number, row in enumerate(rows, start=2):
            if all(is_blank(cell) for cell in row):
                continue
            data.append((row_number, list(row)))
            if max_rows and len(data) > max_rows:
                raise CustomException(
                    ExceptionType.BAD_REQUEST, f"Too many rows: the limit is {max_rows}"
                )
        return header, data
    finally:
        workbook.close()


def xlsx_response(content: bytes, filename: str) -> Response:
    """Small in-memory workbook (e.g. the import template)."""
    return Response(
        content,
        media_type=XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
