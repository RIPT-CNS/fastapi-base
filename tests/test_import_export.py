import io
import json
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

from app.core.database import SessionLocal
from app.utils.excel import XLSX_MEDIA_TYPE
from tests.conftest import bearer, login
from tests.crud_resources import Item, ItemAPI

URL = "/api/test-items"
TEMPLATE_HEADERS = ["id", "name", "sku", "price", "tags", "released_on", "category_name"]


def make_xlsx(rows: list[list]) -> bytes:
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def read_xlsx(content: bytes, sheet: str = "data") -> list[list]:
    workbook = load_workbook(io.BytesIO(content))
    return [list(row) for row in workbook[sheet].iter_rows(values_only=True)]


def upload(client, path: str, content: bytes, mode: str = "insert", headers=None):
    return client.post(
        path,
        params={"mode": mode},
        files={"file": ("data.xlsx", content, XLSX_MEDIA_TYPE)},
        headers=headers,
    )


def errors_of(body: dict) -> set[tuple]:
    errors = body["data"]["errors"] if "data" in body and body["data"] else body["errors"]
    return {(e["row"], e["column"], e["field"]) for e in errors}


# --------------------------------------------------------------- template
def test_template(client, user_headers):
    response = client.get(f"{URL}/import/template", headers=user_headers)
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    assert "attachment" in response.headers["content-disposition"]
    assert read_xlsx(response.content)[0] == TEMPLATE_HEADERS
    guide = {row[0]: row for row in read_xlsx(response.content, "guide")[1:]}
    assert guide["name"][2] == "x"  # required
    assert not guide["price"][2] and guide["price"][3] == "Giá bán (VND)"
    assert guide["released_on"][1].startswith("ngày")


# ------------------------------------------------------------- validation
def test_validate_reports_row_and_column_errors(client, user_headers, items):
    content = make_xlsx(
        [
            ["Tên", "sku", "price", "category_name", "notes"],  # header by title; notes ignored
            ["Mango", "M1", 15, "fruit", "x"],  # row 2: ok, but duplicate sku with row 4
            ["X", "M2", -5, None],  # row 3: name too short, negative price
            ["Kiwi", "M1", 7, "fruit"],  # row 4: duplicate sku in file
            ["Pear", "A1", 7, "fruit"],  # row 5: sku already in DB
            [None, None, None, None],  # empty row: skipped
            ["Plum", 12345, "abc", None],  # row 7: price not a number
            ["Lemon", "L1", 2, "unknown"],  # row 8: unknown category (custom rule)
            ["Lime", "L2", 2, None],  # row 9: ok
        ]
    )
    response = upload(client, f"{URL}/import/validate", content, headers=user_headers)
    assert response.status_code == 200, response.text
    data = response.json()["data"]
    assert (data["total_rows"], data["valid_rows"], data["dry_run"]) == (7, 1, True)
    assert data["ignored_columns"] == ["notes"]
    assert errors_of(response.json()) == {
        (2, "B", "sku"),
        (3, "A", "name"),
        (3, "C", "price"),
        (4, "B", "sku"),
        (5, "B", "sku"),
        (7, "C", "price"),
        (8, "D", "category_name"),
    }
    first = data["errors"][0]
    assert first["message"].startswith("Duplicate value in file (rows 2, 4)")

    # The real import refuses the whole file and writes nothing
    response = upload(client, f"{URL}/import", content, headers=user_headers)
    assert response.status_code == 422
    assert response.json()["message"] == "Import rejected: 7 errors in 6 rows"
    assert len(errors_of(response.json())) == 7
    assert client.get(f"{URL}/count").json()["data"] == 4


def test_import_insert(client, user_headers, categories):
    content = make_xlsx(
        [
            TEMPLATE_HEADERS,
            [None, "Mango", 12345, 15.0, "tropical, yellow", date(2024, 5, 1), "fruit"],
            [None, "Leek", "L1", None, '["green"]', None, "vegetable"],
        ]
    )
    response = upload(client, f"{URL}/import", content, headers=user_headers)
    assert response.status_code == 200, response.text
    assert response.json()["data"]["inserted"] == 2
    assert response.json()["message"] == "Imported: 2 inserted, 0 updated"

    rows = client.get(f"{URL}/list", params={"sort": "id"}).json()["data"]
    mango, leek = rows
    assert mango["sku"] == "12345"  # number cell in a text column
    assert mango["price"] == 15
    assert mango["tags"] == ["tropical", "yellow"]
    assert mango["released_on"] == "2024-05-01"
    assert mango["category_id"] == categories["fruit"]
    assert (leek["price"], leek["tags"]) == (0, ["green"])  # empty cell -> schema default


def test_export_edit_reimport(client, user_headers, items):
    response = client.get(f"{URL}/export/importable", params={"sort": "id"})
    assert response.status_code == 200
    rows = read_xlsx(response.content)
    assert rows[0] == TEMPLATE_HEADERS
    apple = rows[1]
    assert apple[:4] == [1, "Apple", "A1", 10]
    assert json.loads(apple[4]) == ["red", "sweet"]
    assert apple[6] == "fruit"

    apple[3] = 77  # edit the price, keep the rest
    response = upload(client, f"{URL}/import", make_xlsx(rows), mode="update", headers=user_headers)
    assert response.status_code == 200, response.text
    assert response.json()["data"]["updated"] == 4

    data = client.get(f"{URL}/1").json()["data"]
    assert (data["price"], data["tags"], data["released_on"]) == (
        77,
        ["red", "sweet"],
        "2024-01-10",
    )
    assert data["category_id"] == items[0]["category_id"]


def test_import_update_mode_errors(client, user_headers, items):
    content = make_xlsx([["id", "price"], [None, 1], [999, 1], ["abc", 1]])
    response = upload(
        client, f"{URL}/import/validate", content, mode="update", headers=user_headers
    )
    messages = {(e["row"], e["message"]) for e in response.json()["data"]["errors"]}
    assert messages == {(2, "Required in update mode"), (3, "Not found"), (4, "Invalid value")}


def test_import_upsert_partial(client, user_headers, items):
    content = make_xlsx([["id", "name"], [1, "Apple v2"], [None, "Brand new"]])
    response = upload(client, f"{URL}/import", content, mode="upsert", headers=user_headers)
    assert response.status_code == 200, response.text
    assert (response.json()["data"]["inserted"], response.json()["data"]["updated"]) == (1, 1)
    apple = client.get(f"{URL}/1").json()["data"]
    assert (apple["name"], apple["price"], apple["sku"]) == ("Apple v2", 10, "A1")  # partial


def test_users_import_by_username(client, admin_headers, normal_user):
    template = read_xlsx(client.get("/api/users/import/template", headers=admin_headers).content)
    assert template[0][0] != "id"  # import_key = username, already a column
    content = make_xlsx(
        [
            ["username", "email", "password", "full_name", "roles"],
            ["user1", None, None, "Updated Name", None],
            ["newbie", "NewBie@Example.com", "Secret@123", "New Bie", "user, editor"],
        ]
    )
    response = upload(client, "/api/users/import", content, mode="upsert", headers=admin_headers)
    assert response.status_code == 200, response.text
    assert (response.json()["data"]["inserted"], response.json()["data"]["updated"]) == (1, 1)

    tokens = login(client, "newbie", "Secret@123")
    me = client.get("/api/auth/me", headers=bearer(tokens["access_token"])).json()["data"]
    assert (me["email"], sorted(me["roles"])) == ("newbie@example.com", ["editor", "user"])
    assert login(client, "user1")["access_token"]  # password untouched by the update

    # Same email (any case) again -> unique error on the email column
    duplicate = make_xlsx(
        [["username", "email", "password"], ["other", "NEWBIE@example.com", "x" * 8]]
    )
    response = upload(client, "/api/users/import/validate", duplicate, headers=admin_headers)
    assert errors_of(response.json()) == {(2, "B", "email")}


def test_import_permissions(client, user_headers):
    content = make_xlsx([["username", "password"], ["x1", "Secret@123"]])
    assert upload(client, "/api/users/import", content, headers=user_headers).status_code == 403
    assert client.get("/api/users/export", headers=user_headers).status_code == 403


# ----------------------------------------------------------------- export
def test_export_fields_and_relations(client, items):
    response = client.get(
        f"{URL}/export",
        params={"fields": "name,category.name,tags,category", "sort": "name", "price__gte": 5},
    )
    assert response.status_code == 200
    rows = read_xlsx(response.content)
    assert rows[0] == ["name", "category.name", "tags", "category"]
    assert [row[0] for row in rows[1:]] == ["100% Juice", "Apple", "Banana"]
    juice, apple = rows[1], rows[2]
    assert juice[1:] == [None, "[]", None]
    assert apple[1] == "fruit"
    assert json.loads(apple[2]) == ["red", "sweet"]
    assert json.loads(apple[3])["name"] == "fruit"  # to-one relation -> JSON object

    # to-many relation -> JSON array
    response = client.get("/api/test-categories/export", params={"fields": "name,items.name,items"})
    fruit = next(row for row in read_xlsx(response.content) if row[0] == "fruit")
    assert sorted(json.loads(fruit[1])) == ["Apple", "Banana"]
    assert {item["sku"] for item in json.loads(fruit[2])} == {"A1", "B1"}

    default = read_xlsx(client.get(f"{URL}/export").content)[0]
    assert default[:3] == ["id", "name", "sku"] and "deleted_at" in default

    assert client.get(f"{URL}/export", params={"fields": "nope"}).status_code == 400


def test_export_does_not_write_formulas(client, user_headers):
    client.post(URL, json={"name": '=HYPERLINK("http://evil","x")'}, headers=user_headers)
    content = client.get(f"{URL}/export", params={"fields": "name"}).content
    cell = load_workbook(io.BytesIO(content))["data"]["A2"]
    assert (cell.data_type, cell.value) == ("s", '=HYPERLINK("http://evil","x")')


# ------------------------------------------------------------ bad uploads
def test_bad_uploads(client, user_headers, monkeypatch):
    path = f"{URL}/import/validate"
    assert upload(client, path, b"not an excel file", headers=user_headers).status_code == 400
    assert upload(client, path, b"", headers=user_headers).status_code == 400
    assert upload(client, path, make_xlsx([[None]]), headers=user_headers).status_code == 400

    monkeypatch.setattr(ItemAPI, "import_max_rows", 1)
    content = make_xlsx([["name"], ["aa"], ["bb"]])
    response = upload(client, path, content, headers=user_headers)
    assert (response.status_code, response.json()["message"]) == (
        400,
        "Too many rows: the limit is 1",
    )

    monkeypatch.setattr(ItemAPI, "import_max_file_size_mb", 0)
    assert upload(client, path, content, headers=user_headers).status_code == 413


def test_nothing_saved_when_save_fails(client, user_headers, monkeypatch):
    def boom(self, row, ctx):
        if row.row == 3:
            raise RuntimeError("boom")
        return type(self).__mro__[1].import_save_row(self, row, ctx)

    from tests.crud_resources import ItemService

    monkeypatch.setattr(ItemService, "import_save_row", boom)
    content = make_xlsx([["name"], ["aa"], ["bb"]])
    with pytest.raises(RuntimeError):
        upload(client, f"{URL}/import", content, headers=user_headers)
    with SessionLocal() as db:
        assert db.query(Item).count() == 0
